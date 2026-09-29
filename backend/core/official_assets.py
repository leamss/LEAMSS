import io
from datetime import datetime
from typing import Optional
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import qrcode
from PIL import Image, ImageDraw, ImageFont
import docx

def docx_to_pdf(docx_bytes: bytes, client_name: str = "Valued Applicant") -> bytes:
    """Convert Microsoft Word DOCX bytes into a genuine ReportLab PDF."""
    try:
        doc_in = docx.Document(io.BytesIO(docx_bytes))
    except Exception:
        return generate_official_sla_pdf(client_name=client_name)

    buf = io.BytesIO()
    doc_out = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()
    teal = colors.HexColor("#12433B")
    gold = colors.HexColor("#C99A3B")
    slate = colors.HexColor("#334155")
    dark = colors.HexColor("#0F172A")
    gray_bg = colors.HexColor("#F8FAFC")
    border_color = colors.HexColor("#CBD5E1")

    h1_style = ParagraphStyle(
        "DocxH1",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=15,
        textColor=teal,
        spaceBefore=10,
        spaceAfter=4,
    )

    h2_style = ParagraphStyle(
        "DocxH2",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9.5,
        leading=13,
        textColor=slate,
        spaceBefore=6,
        spaceAfter=2,
    )

    body_style = ParagraphStyle(
        "DocxBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=12,
        textColor=dark,
        spaceAfter=4,
    )

    bullet_style = ParagraphStyle(
        "DocxBullet",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=12,
        textColor=slate,
        leftIndent=12,
        firstLineIndent=-8,
        spaceAfter=3,
    )

    story = []

    # Header
    header_data = [
        [
            Paragraph("<b>LADHANI EDUCATION & MIGRATION SERVICES PVT. LTD.</b>", ParagraphStyle("Hdr1", fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=colors.white)),
            Paragraph("<b>SERVICE LEVEL AGREEMENT (SLA)</b><br/>Global Immigration Division", ParagraphStyle("Hdr2", fontName="Helvetica", fontSize=8, leading=10, textColor=gold, alignment=2))
        ]
    ]
    t_hdr = Table(header_data, colWidths=[350, 165])
    t_hdr.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), teal),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("PADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(t_hdr)
    story.append(Spacer(1, 10))

    # Add paragraphs from DOCX
    for p in doc_in.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        
        # Replace empty client name if present
        if txt.startswith("Name:") and len(txt) < 10 and client_name:
            txt = f"Name: {client_name}"
        
        # Heading detection
        if any(txt.startswith(prefix) for prefix in ["CLIENT DETAILS", "ANNEXURE", "Annexure", "1.", "2.", "3.", "4.", "5.", "6.", "7.", "Stage 1", "Stage 2", "Stage 3", "LEAMSS Services", "Working together"]):
            story.append(Paragraph(f"<b>{txt}</b>", h1_style))
        elif txt.startswith("•") or txt.startswith("-") or txt.startswith("*"):
            clean_txt = txt.lstrip("•-* ").strip()
            story.append(Paragraph(f"• {clean_txt}", bullet_style))
        else:
            is_bold = any(run.bold for run in p.runs) if p.runs else False
            if is_bold and len(txt) < 80:
                story.append(Paragraph(f"<b>{txt}</b>", h2_style))
            else:
                story.append(Paragraph(txt, body_style))

    # Add tables from DOCX
    for t in doc_in.tables:
        t_data = []
        for row in t.rows:
            row_cells = []
            for cell in row.cells:
                cell_txt = cell.text.strip()
                row_cells.append(Paragraph(cell_txt or "-", body_style))
            if row_cells:
                t_data.append(row_cells)
        if t_data:
            num_cols = max(len(r) for r in t_data)
            col_w = 515 / num_cols
            t_elem = Table(t_data, colWidths=[col_w] * num_cols)
            t_elem.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), gray_bg),
                ("BOX", (0, 0), (-1, -1), 1, border_color),
                ("INNERGRID", (0, 0), (-1, -1), 0.5, border_color),
                ("PADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(Spacer(1, 6))
            story.append(t_elem)
            story.append(Spacer(1, 6))

    doc_out.build(story)
    return buf.getvalue()


def generate_official_sla_pdf(client_name: str = "Valued Applicant", country: str = "Australia") -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    
    teal = colors.HexColor("#12433B")
    teal_light = colors.HexColor("#EBF3F1")
    gold = colors.HexColor("#C99A3B")
    slate = colors.HexColor("#334155")
    dark = colors.HexColor("#0F172A")
    gray_bg = colors.HexColor("#F8FAFC")
    border_color = colors.HexColor("#CBD5E1")

    title_style = ParagraphStyle(
        "SlaTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=teal,
        alignment=1,
    )
    
    subtitle_style = ParagraphStyle(
        "SlaSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=gold,
        alignment=1,
    )

    h1_style = ParagraphStyle(
        "SlaH1",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=16,
        textColor=teal,
        spaceBefore=10,
        spaceAfter=4,
    )

    body_style = ParagraphStyle(
        "SlaBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=12.5,
        textColor=dark,
    )

    bullet_style = ParagraphStyle(
        "SlaBullet",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=12,
        textColor=slate,
        leftIndent=12,
        firstLineIndent=-8,
    )

    table_header = ParagraphStyle(
        "SlaTH",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11,
        textColor=colors.white,
    )

    table_cell = ParagraphStyle(
        "SlaTC",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        leading=11,
        textColor=dark,
    )

    table_cell_bold = ParagraphStyle(
        "SlaTCB",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=11,
        textColor=teal,
    )

    story = []

    # ── HEADER BANNER ──
    header_data = [
        [
            Paragraph("<b>LADHANI EDUCATION & MIGRATION SERVICES PVT. LTD.</b>", ParagraphStyle("Hdr1", fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=colors.white)),
            Paragraph("<b>GLOBAL IMMIGRATION EXPERTS</b><br/>ISO 9001:2015 Certified", ParagraphStyle("Hdr2", fontName="Helvetica", fontSize=8.5, leading=11, textColor=gold, alignment=2))
        ]
    ]
    t_hdr = Table(header_data, colWidths=[360, 160])
    t_hdr.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), teal),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("PADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
    ]))
    story.append(t_hdr)
    story.append(Spacer(1, 10))

    # ── TITLE & METADATA ──
    story.append(Paragraph("SERVICE LEVEL AGREEMENT (SLA) & RETAINER CONTRACT", title_style))
    story.append(Paragraph("Permanent Residency & Skilled Migration Consulting Services", subtitle_style))
    story.append(Spacer(1, 8))

    meta_table_data = [
        [
            Paragraph("<b>Client Name:</b>", table_cell_bold),
            Paragraph(client_name or "Applicant", table_cell),
            Paragraph("<b>Country & Pathway:</b>", table_cell_bold),
            Paragraph(f"{country} (Subclasses 189 / 190 / 491)", table_cell),
        ],
        [
            Paragraph("<b>Agreement Date:</b>", table_cell_bold),
            Paragraph(datetime.now().strftime("%d %B %Y"), table_cell),
            Paragraph("<b>Consulting Body:</b>", table_cell_bold),
            Paragraph("LEAMSS Migration Division", table_cell),
        ]
    ]
    t_meta = Table(meta_table_data, colWidths=[95, 165, 110, 150])
    t_meta.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), teal_light),
        ("BOX", (0, 0), (-1, -1), 1, teal),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, border_color),
        ("PADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 10))

    # ── 1. SCOPE OF SERVICES ──
    story.append(Paragraph("1. Scope of Services & Deliverables", h1_style))
    story.append(Paragraph(
        "Ladhani Education & Migration Services Pvt. Ltd. ('LEAMSS') agrees to provide end-to-end professional migration consulting, profile evaluation, skills assessment documentation, Expression of Interest (EOI) lodgement, state nomination filing, and visa filing assistance for the Client. Deliverables include:",
        body_style
    ))
    story.append(Spacer(1, 4))
    story.append(Paragraph("• <b>ANZSCO Code & Pathway Advisory:</b> Detailed mapping of qualifications and work experience to relevant assessing authorities.", bullet_style))
    story.append(Paragraph("• <b>Skills Assessment Dossier Preparation:</b> Complete drafting, verification of reference letters, certified documentation, and lodgement with assessing authorities (e.g. ACS, VETASSESS, Engineers Australia, TRA).", bullet_style))
    story.append(Paragraph("• <b>SkillSelect EOI & State Nomination:</b> Strategic submission of Expression of Interest (EOI) and State/Territory Nomination matrices across high-invitation states.", bullet_style))
    story.append(Paragraph("• <b>Department of Home Affairs (DHA) Visa Filing:</b> End-to-end visa dossier compilation, Form 80/1221 preparation, police clearance and medicals guidance.", bullet_style))
    story.append(Paragraph("• <b>Dedicated Case Officer Support:</b> One-on-one assistance by an assigned senior case manager with regular bi-weekly status updates.", bullet_style))

    story.append(Spacer(1, 8))

    # ── 2. STAGES & PROCESSING TIMELINE ──
    story.append(Paragraph("2. Three-Stage Application Procedure", h1_style))
    
    stage_table_data = [
        [
            Paragraph("Stage", table_header),
            Paragraph("Milestone & Scope", table_header),
            Paragraph("Key Responsibilities", table_header),
        ],
        [
            Paragraph("<b>Stage 1</b>", table_cell_bold),
            Paragraph("<b>Skills Assessment Lodgement</b><br/>Filing of qualification and employment evaluation with designated authority.", table_cell),
            Paragraph("Client provides verified documents & English score (IELTS/PTE). LEAMSS compiles technical dossier & submits.", table_cell),
        ],
        [
            Paragraph("<b>Stage 2</b>", table_cell_bold),
            Paragraph("<b>EOI & State Nomination Pool</b><br/>SkillSelect pool entry and state nomination invitation tracking.", table_cell),
            Paragraph("LEAMSS files EOI & manages state applications. Application remains active in the pool for up to 2 years.", table_cell),
        ],
        [
            Paragraph("<b>Stage 3</b>", table_cell_bold),
            Paragraph("<b>DHA Visa Lodgement & Grant</b><br/>Formal visa application submission within 60 days of invitation receipt.", table_cell),
            Paragraph("LEAMSS reviews all supporting evidence, lodges visa with DHA, and manages case correspondence till decision.", table_cell),
        ]
    ]
    t_stage = Table(stage_table_data, colWidths=[55, 235, 230])
    t_stage.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), teal),
        ("BACKGROUND", (0, 1), (-1, -1), gray_bg),
        ("BOX", (0, 0), (-1, -1), 1, border_color),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, border_color),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("PADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(t_stage)

    story.append(Spacer(1, 10))

    # ── 3. REFUND POLICY & GUARANTEE ──
    story.append(Paragraph("3. 100% Performance Refund Policy & Commitments", h1_style))
    story.append(Paragraph(
        "LEAMSS stands firmly behind our quality standards. Our professional fee is covered by a <b>100% Refund Policy</b> under the following terms:",
        body_style
    ))
    story.append(Spacer(1, 4))
    story.append(Paragraph("• <b>Negative Skills Assessment:</b> 100% refund of professional fees if skills assessment is refused on technical grounds (excluding document forgery or negative employer verification).", bullet_style))
    story.append(Paragraph("• <b>Visa Refusal:</b> 100% refund of professional fees if visa application is refused due to any administrative error on part of LEAMSS.", bullet_style))
    story.append(Paragraph("• <b>Non-Refundable Scenarios:</b> Voluntary withdrawal by applicant after receiving document checklist, failure to achieve minimum required English test score, submission of false documents, or rejection due to medical/character grounds.", bullet_style))

    story.append(Spacer(1, 10))

    # ── 4. POST-LANDING & SETTLEMENT ASSISTANCE ──
    story.append(Paragraph("4. Post-Landing & Value-Added Support Services", h1_style))
    story.append(Paragraph(
        "Upon visa grant, LEAMSS provides complete settlement support to ensure a seamless transition:",
        body_style
    ))
    story.append(Spacer(1, 3))
    story.append(Paragraph("• Australian bank account opening assistance prior to departure.", bullet_style))
    story.append(Paragraph("• Medicare registration and Tax File Number (TFN) application assistance.", bullet_style))
    story.append(Paragraph("• CV / Resume Australian format optimization and distribution across top tier recruitment networks.", bullet_style))
    story.append(Paragraph("• 1-hour 1-on-1 Employment Readiness & Orientation session with Australian resident advisor.", bullet_style))

    story.append(Spacer(1, 12))

    # ── 5. SIGNATURE & ACCEPTANCE ──
    story.append(Paragraph("5. Authorization & Signatures", h1_style))
    story.append(Paragraph(
        "By proceeding with fee payment or electronically accepting this Service Level Agreement, the Client agrees to the terms and conditions outlined above. This agreement is legally binding under applicable IT Act and international contract standards.",
        body_style
    ))
    story.append(Spacer(1, 10))

    sig_table_data = [
        [
            Paragraph("<b>FOR THE CLIENT / APPLICANT:</b>", table_cell_bold),
            Paragraph("<b>FOR LADHANI EDUCATION & MIGRATION SERVICES:</b>", table_cell_bold),
        ],
        [
            Paragraph(f"<b>Name:</b> {client_name or 'Applicant'}<br/><b>Date:</b> {datetime.now().strftime('%d %B %Y')}<br/><b>Status:</b> Accepted Electronically", table_cell),
            Paragraph("<b>Authorized Signatory:</b> LEAMSS Migration Division<br/><b>Company:</b> Ladhani Education & Migration Services Pvt. Ltd.<br/><b>Contact:</b> info@leamss.com | +91 77188 82427", table_cell),
        ]
    ]
    t_sig = Table(sig_table_data, colWidths=[260, 260])
    t_sig.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), gray_bg),
        ("BOX", (0, 0), (-1, -1), 1, teal),
        ("INNERGRID", (0, 0), (-1, -1), 0.5, border_color),
        ("PADDING", (0, 0), (-1, -1), 8),
    ]))
    story.append(t_sig)

    doc.build(story)
    return buf.getvalue()


def ensure_valid_sla_pdf(raw_data: Optional[bytes] = None, client_name: str = "Valued Applicant", country: str = "Australia") -> bytes:
    """Ensure SLA data is 100% valid PDF bytes (starting with %PDF). If DOCX, converts to PDF; if none, generates official PDF."""
    if raw_data:
        if raw_data.startswith(b"%PDF"):
            return raw_data
        if raw_data.startswith(b"PK"):
            try:
                converted = docx_to_pdf(raw_data, client_name=client_name)
                if converted and converted.startswith(b"%PDF"):
                    return converted
            except Exception:
                pass
    return generate_official_sla_pdf(client_name=client_name, country=country)


def generate_official_qr_image(amount_str: str = "80,000", upi_id: str = "7738352427@okbizaxis") -> bytes:
    upi_uri = f"upi://pay?pa={upi_id}&pn=Ladhani%20Education%20and%20Migration%20Services&cu=INR"
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=8,
        border=2,
    )
    qr.add_data(upi_uri)
    qr.make(fit=True)
    qr_img = qr.make_image(fill_color="#12433B", back_color="white").convert("RGB")

    card_w, card_h = 480, 560
    card = Image.new("RGB", (card_w, card_h), "#F8FAFC")
    draw = ImageDraw.Draw(card)

    draw.rectangle([(0, 0), (card_w, 80)], fill="#12433B")
    draw.text((card_w // 2, 28), "LEAMSS OFFICIAL PAYMENT QR", fill="#FFFFFF", anchor="mm")
    draw.text((card_w // 2, 54), "Ladhani Education & Migration Services Pvt. Ltd.", fill="#C99A3B", anchor="mm")

    qr_x = (card_w - qr_img.width) // 2
    qr_y = 100
    draw.rectangle([(qr_x - 6, qr_y - 6), (qr_x + qr_img.width + 6, qr_y + qr_img.height + 6)], fill="#FFFFFF", outline="#CBD5E1", width=2)
    card.paste(qr_img, (qr_x, qr_y))

    text_y = qr_y + qr_img.height + 20
    draw.text((card_w // 2, text_y), f"UPI ID: {upi_id}", fill="#12433B", anchor="mm")
    draw.text((card_w // 2, text_y + 24), "Accepted: Google Pay · PhonePe · Paytm · BHIM · Axis Bank", fill="#64748B", anchor="mm")
    
    box_top = text_y + 44
    draw.rectangle([(25, box_top), (card_w - 25, box_top + 100)], fill="#FFFFFF", outline="#12433B", width=1)
    draw.text((38, box_top + 12), "Bank: AXIS Bank | Dombivali East Branch", fill="#0F172A")
    draw.text((38, box_top + 32), "A/C Name: Ladhani Education & Migration Services OPC Pvt Ltd", fill="#0F172A")
    draw.text((38, box_top + 52), "Current A/C No: 924020021437590 | IFSC: UTIB0001364", fill="#12433B")
    draw.text((38, box_top + 74), "Razorpay Link: https://rzp.io/rzp/IndepdenceJjMJwx1", fill="#D4633F")

    buf = io.BytesIO()
    card.save(buf, format="PNG")
    return buf.getvalue()
